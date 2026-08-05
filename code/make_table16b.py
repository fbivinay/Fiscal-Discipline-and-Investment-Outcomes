# -*- coding: utf-8 -*-
"""Publish the substituted ranking that Section 4.7 currently only describes.

The paper's own criterion is that a ranking should be shown against an outcome
rather than asserted. Section 4.7 reports that substituting the revenue deficit
for derived capital expenditure moves the table, and names four of the moves,
but never prints the result. This adds it as Table 16b, computed from the same
five-year averages as Table 15.
"""
import copy, sys
import pandas as pd
from scipy import stats
import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = r'paper\Fiscal_Discipline_and_Investment_Outcomes.docx'

d = pd.read_csv('data/processed/MASTER_PANEL_DATASET.csv')
d.columns = ["state", "year", "year_num", "fd", "fd_amt", "debt", "debt_amt", "otr",
             "otr_amt", "capex", "capex_amt", "rd", "rd_amt", "fdi", "gsdp", "growth",
             "log_fdi", "log_gsdp", "covid", "post_covid", "literacy", "urban",
             "rank_in_year"]
g = d.groupby("state").agg(fd=("fd", "mean"), debt=("debt", "mean"), otr=("otr", "mean"),
                           capex=("capex", "mean"), rd=("rd", "mean"), fdi=("fdi", "sum"))
published = ((g.fd.rank() + g.debt.rank() + g.otr.rank(ascending=False)
              + g.capex.rank(ascending=False)) / 4).rank()
substituted = ((g.fd.rank() + g.debt.rank() + g.otr.rank(ascending=False)
                + g.rd.rank()) / 4).rank()
debt_only = g.debt.rank()
fdi_rank = g.fdi.rank(ascending=False)

# the prose in Section 4.7 must still be true of what we are about to print
assert round(stats.spearmanr(published, substituted)[0], 3) == 0.930
assert substituted.idxmin() == "Maharashtra" and substituted["Uttar Pradesh"] == 2.0
assert published["Telangana"] == 3.0 and substituted["Telangana"] == 5.5
assert substituted["West Bengal"] < substituted["Rajasthan"]

t = pd.DataFrame({"published": published, "substituted": substituted,
                  "debt_only": debt_only, "debt": g.debt, "fdi": fdi_rank}
                 ).sort_values("published")


def fmt(x):
    return str(int(x)) if float(x).is_integer() else f"{x:.1f}"


doc = docx.Document(SRC)
body = doc.element.body
kids = list(body.iterchildren())

# anchor on Table 16a's own table so the new one lands after it, inside 4.7
anchor = None
for i, ch in enumerate(kids):
    if ch.tag == qn("w:p") and Paragraph(ch, doc).text.strip().startswith("Table 16a."):
        anchor = kids[i + 1]
        assert anchor.tag == qn("w:tbl"), "Table 16a is not followed by a table"
        break
if anchor is None:
    sys.exit("Table 16a caption not found")

caption_src = None
for ch in kids:
    if ch.tag == qn("w:p") and Paragraph(ch, doc).text.strip().startswith("Table 16a."):
        caption_src = ch
        break

HEAD = ["State", "Published composite rank", "Revenue-deficit substitution",
        "Debt-to-GSDP alone", "Debt (% GSDP)", "Realised FDI rank"]
tbl = doc.add_table(rows=len(t) + 1, cols=len(HEAD))
tbl.style = doc.tables[0].style
for j, h in enumerate(HEAD):
    tbl.rows[0].cells[j].paragraphs[0].add_run(h)
for i, (state, r) in enumerate(t.iterrows(), start=1):
    vals = [state, fmt(r.published), fmt(r.substituted), fmt(r.debt_only),
            f"{r.debt:.2f}", fmt(r.fdi)]
    for j, v in enumerate(vals):
        tbl.rows[i].cells[j].paragraphs[0].add_run(v)

cap = copy.deepcopy(caption_src)
cap_par = Paragraph(cap, doc)
cap_par.runs[0].text = (
    "Table 16b. The composite ranking under the redundancy fix, against the single "
    "indicator and against the outcome. Column 2 is the ranking published in Table 15; "
    "column 3 replaces derived capital expenditure with the revenue deficit, which removes "
    "the double-counting of the fiscal deficit described above; column 4 ranks on debt alone. "
    "The two composites correlate at rho = 0.930. Against realised five-year FDI the "
    "published composite gives rho = -0.255 (p = 0.477), the substituted one -0.433 "
    "(p = 0.211) and debt alone -0.842 (p = 0.002).")
for r in cap_par.runs[1:]:
    r.text = ''

anchor.addnext(tbl._tbl)
anchor.addnext(cap)
doc.save(SRC)
print("Table 16b inserted")
print(t.to_string())
